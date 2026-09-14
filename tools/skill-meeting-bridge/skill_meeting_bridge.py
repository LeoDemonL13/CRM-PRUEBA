import os, io, re, json, time, wave, smtplib, asyncio, tempfile, subprocess, base64
from pathlib import Path
from collections import deque
from email.message import EmailMessage
from fastapi import FastAPI, WebSocket, WebSocketDisconnect, HTTPException, UploadFile, File, Form, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

SAMPLE_RATE=int(os.getenv('SKILL_SAMPLE_RATE','16000'))
TOKEN=os.getenv('SKILL_BRIDGE_TOKEN','CAMBIA_ESTE_TOKEN')
OUT=Path(os.getenv('SKILL_OUTPUT_DIR','./reuniones')).resolve(); OUT.mkdir(parents=True,exist_ok=True)
KEEP_AUDIO=os.getenv('SKILL_KEEP_AUDIO','false').lower() in ('1','true','yes','si','sí')
WHISPER_MODEL=os.getenv('SKILL_WHISPER_MODEL','tiny')
WHISPER_FINAL_MODEL=os.getenv('SKILL_WHISPER_FINAL_MODEL','small')
LANGUAGE=os.getenv('SKILL_LANGUAGE','es')
OLLAMA_URL=os.getenv('OLLAMA_BASE_URL','http://127.0.0.1:11434').rstrip('/')
OLLAMA_MODEL=os.getenv('SKILL_OLLAMA_MODEL','qwen3:4b')
USE_OLLAMA=os.getenv('SKILL_USE_OLLAMA_SUMMARY','true').lower() in ('1','true','yes','si','sí')
HF_TOKEN=os.getenv('HF_TOKEN','')
PYANNOTE_MODEL=os.getenv('PYANNOTE_MODEL_PATH','pyannote/speaker-diarization-community-1')
START_PATTERNS=[r'\bskill\b.*\b(?:inicia|iniciar|empezamos|empieza|comienza|comenzar)\b.*\b(?:reunion|junta|grabacion)\b',r'\b(?:vamos a iniciar|vamos a empezar)\b.*\b(?:reunion|junta)\b']
END_PATTERNS=[r'\b(?:ya termino|ya termino la|ya terminamos|terminamos|finaliza|finalizar|acaba|acabamos|cerrar)\b.*\b(?:reunion|junta|grabacion)\b.*\bskill\b',r'\bskill\b.*\b(?:termina|finaliza|cierra)\b.*\b(?:reunion|junta|grabacion)\b',r'\bya termino la reunion skill gracias\b']
app=FastAPI(title='SKILL Meeting Bridge',version='136')
app.add_middleware(CORSMiddleware,allow_origins=['*'],allow_credentials=False,allow_methods=['*'],allow_headers=['*'])
clients={}
whisper_cache={}
pyannote_pipeline=None

class EmailRequest(BaseModel):
    to:list[str]
    subject:str='Minuta de reunión SKILL'
    body:str='Adjunto encontrarás la minuta de reunión generada por SKILL.'
    docx:str=''
    pdf:str=''

class BrowserAttachment(BaseModel):
    filename:str
    content:str
    contentType:str='application/octet-stream'

class BrowserEmailRequest(BaseModel):
    to:list[str]
    subject:str='Minuta de reunión SKILL'
    body:str='Adjunto encontrarás la minuta de reunión generada por SKILL.'
    attachments:list[BrowserAttachment]=Field(default_factory=list)

def norm(v):
    import unicodedata
    s=unicodedata.normalize('NFD',str(v or '')).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9ñ\s]',' ',s).replace('  ',' ').strip()

def write_wav(path,pcm):
    with wave.open(str(path),'wb') as wf:
        wf.setnchannels(1); wf.setsampwidth(2); wf.setframerate(SAMPLE_RATE); wf.writeframes(bytes(pcm))

def whisper(model_name):
    if model_name not in whisper_cache:
        from faster_whisper import WhisperModel
        compute=os.getenv('SKILL_WHISPER_COMPUTE','int8')
        whisper_cache[model_name]=WhisperModel(model_name,device='cpu',compute_type=compute)
    return whisper_cache[model_name]

def transcribe_pcm(pcm,model_name=None,words=False):
    if not pcm:return []
    model=whisper(model_name or WHISPER_MODEL)
    with tempfile.NamedTemporaryFile(suffix='.wav',delete=False) as tmp:
        path=Path(tmp.name)
    try:
        write_wav(path,pcm)
        segments,_=model.transcribe(str(path),language=LANGUAGE,vad_filter=True,beam_size=2,word_timestamps=words)
        return list(segments)
    finally:
        path.unlink(missing_ok=True)

def transcribe_file(path,model_name=None,words=False):
    model=whisper(model_name or WHISPER_FINAL_MODEL)
    kwargs=dict(language=LANGUAGE,vad_filter=True,beam_size=5,word_timestamps=words,condition_on_previous_text=True)
    try:
        segments,_=model.transcribe(str(path),vad_parameters={'min_silence_duration_ms':350,'speech_pad_ms':220},**kwargs)
    except TypeError:
        segments,_=model.transcribe(str(path),**kwargs)
    return list(segments)

def duration_seconds(path):
    try:
        import av
        container=av.open(str(path))
        if container.duration:return float(container.duration/av.time_base)
        stream=next((x for x in container.streams if x.type=='audio'),None)
        if stream and stream.duration and stream.time_base:return float(stream.duration*stream.time_base)
    except Exception:pass
    return 0.0

def finalize_uploaded_audio(path,title='Reunión de trabajo',participants='browser'):
    diar=diarize(path)
    segs=transcribe_file(path,WHISPER_FINAL_MODEL,True)
    rows=rows_from_segments(segs,diar)
    summary_rows=[{**r,'speaker':r.get('speaker') or 'Voz 1'} for r in rows]
    summary=ollama_summary(summary_rows)
    voices=len(set(r['speaker'] for r in rows if r.get('speaker'))) if rows else 0
    duration=duration_seconds(path)
    meta={'title':title or 'Reunión de trabajo','date':time.strftime('%d/%m/%Y %H:%M'),'duration':time.strftime('%H:%M:%S',time.gmtime(duration)),'speakers':voices,'device':participants or 'browser','diarization':bool(diar),'whisperModel':WHISPER_FINAL_MODEL,'pyannoteModel':PYANNOTE_MODEL if diar else ''}
    stamp=time.strftime('%Y%m%d-%H%M%S');slug='browser'
    base=OUT/f'{stamp}_{slug}'
    docx=Path(str(base)+'.docx');pdf=Path(str(base)+'.pdf');js=Path(str(base)+'.json')
    build_docx(meta,summary_rows,summary,docx);build_pdf(meta,summary_rows,summary,pdf)
    payload={'meta':meta,'summary':summary,'transcript':rows,'files':{'docx':docx.name,'pdf':pdf.name,'json':js.name}}
    js.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    return payload

def command_text(pcm):
    segs=transcribe_pcm(pcm,WHISPER_MODEL,False)
    return ' '.join((s.text or '').strip() for s in segs).strip()

def is_start(text):
    n=norm(text); return any(re.search(p,n) for p in START_PATTERNS)

def is_end(text):
    n=norm(text); return any(re.search(p,n) for p in END_PATTERNS)

def diarize(path):
    global pyannote_pipeline
    if not HF_TOKEN:return []
    try:
        if pyannote_pipeline is None:
            from pyannote.audio import Pipeline
            pyannote_pipeline=Pipeline.from_pretrained(PYANNOTE_MODEL,token=HF_TOKEN)
        result=pyannote_pipeline(str(path))
        annotation=getattr(result,'exclusive_speaker_diarization',None)
        if annotation is None:annotation=getattr(result,'speaker_diarization',None)
        if annotation is None:annotation=result
        rows=[]
        for turn,_,speaker in annotation.itertracks(yield_label=True):rows.append({'start':float(turn.start),'end':float(turn.end),'speaker':str(speaker)})
        return rows
    except Exception as exc:
        print('pyannote:',exc);return []

def speaker_at(start,end,diar):
    if not diar:return ''
    start=float(start or 0); end=max(start,float(end or start))
    labels=[]
    for d in diar:
        if d['speaker'] not in labels:labels.append(d['speaker'])
    best=None; best_overlap=0.0
    for d in diar:
        overlap=max(0.0,min(end,d['end'])-max(start,d['start']))
        if overlap>best_overlap:
            best_overlap=overlap; best=d
    if best is None:
        mid=(start+end)/2
        best=min(diar,key=lambda d:min(abs(mid-d['start']),abs(mid-d['end'])))
    return f'Voz {labels.index(best["speaker"])+1}' if best and best['speaker'] in labels else 'Voz 1'

def rows_from_segments(segs,diar):
    rows=[]
    for seg in segs:
        words=list(getattr(seg,'words',None) or [])
        if words:
            for word in words:
                start=float(word.start if word.start is not None else seg.start)
                end=float(word.end if word.end is not None else seg.end)
                token=str(word.word or '').strip()
                if not token:continue
                label=speaker_at(start,end,diar)
                if rows and rows[-1]['speaker']==label and start-rows[-1]['end']<1.15:
                    rows[-1]['text']=(rows[-1]['text']+' '+token).strip()
                    rows[-1]['end']=end
                else:
                    rows.append({'start':start,'end':end,'time':time.strftime('%H:%M:%S',time.gmtime(start)),'speaker':label,'text':token})
        else:
            start=float(seg.start);end=float(seg.end);token=(seg.text or '').strip()
            if not token:continue
            label=speaker_at(start,end,diar)
            rows.append({'start':start,'end':end,'time':time.strftime('%H:%M:%S',time.gmtime(start)),'speaker':label,'text':token})
    cleaned=[]
    for row in rows:
        token=re.sub(r'\s+',' ',str(row.get('text') or '')).strip()
        if not token:continue
        row={**row,'text':token}
        if cleaned:
            prev=cleaned[-1]
            a=norm(prev.get('text'));b=norm(token)
            if a and b and prev.get('speaker')==row.get('speaker') and float(row.get('start',0))-float(prev.get('end',0))<1.35:
                if b==a or (len(b)>24 and b in a):
                    prev['end']=max(float(prev.get('end',0)),float(row.get('end',0)));continue
                if len(a)>24 and a in b:
                    prev.update(row);continue
                prev['text']=(prev['text']+' '+token).strip();prev['end']=row.get('end',prev.get('end'));continue
        recent=cleaned[-6:]
        if any(norm(r.get('text'))==norm(token) and abs(float(row.get('start',0))-float(r.get('start',0)))<8 for r in recent):continue
        cleaned.append(row)
    return cleaned

def heuristic_summary(rows):
    text=' '.join(r['text'] for r in rows)
    lines=[x.strip() for x in re.split(r'(?<=[.!?])\s+',text) if len(x.strip())>15]
    agreements=[x for x in lines if re.search(r'\b(acord|queda|vamos a|se va a|debe|pendiente|responsable|compromiso|fecha limite)\b',norm(x))]
    questions=[x for x in lines if '?' in x or re.search(r'\b(quien|que|como|cuando|donde|cuanto)\b',norm(x))]
    return {'resumen':' '.join(lines[:6])[:1800] or 'Sin resumen automático.','decisiones':agreements[:12],'pendientes':agreements[:12],'preguntas':questions[:12]}

def ollama_summary(rows):
    if not USE_OLLAMA:return heuristic_summary(rows)
    import urllib.request
    transcript='\n'.join(f'[{r["speaker"]}] {r["text"]}' for r in rows)[:24000]
    system='''Genera una minuta profesional en JSON en español. Devuelve: resumen, temas, decisiones, acuerdos, pendientes, responsables, fechas, preguntas. No inventes. Si un responsable o fecha no fue dicho, usa cadena vacía. acuerdos y pendientes pueden ser listas de objetos con texto, responsable y fecha.'''
    body={'model':OLLAMA_MODEL,'stream':False,'format':'json','messages':[{'role':'system','content':system},{'role':'user','content':transcript}]}
    req=urllib.request.Request(f'{OLLAMA_URL}/api/chat',data=json.dumps(body).encode(),headers={'Content-Type':'application/json'},method='POST')
    try:
        with urllib.request.urlopen(req,timeout=90) as response:data=json.loads(response.read().decode())
        return json.loads(data.get('message',{}).get('content','{}'))
    except Exception:
        return heuristic_summary(rows)

def build_docx(meta,rows,summary,path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    p=d.add_paragraph(); r=p.add_run('SKILLED PROYECTOS INDUSTRIALES'); r.bold=True; r.font.color.rgb=RGBColor(0,65,107); r.font.size=Pt(10)
    h=d.add_heading(meta['title'],0); h.runs[0].font.color.rgb=RGBColor(0,65,107)
    d.add_paragraph(f"Fecha: {meta['date']}\nDuración: {meta['duration']}\nVoces diferenciadas: {meta['speakers']}")
    d.add_heading('Resumen ejecutivo',level=1); d.add_paragraph(str(summary.get('resumen') or 'Sin resumen.'))
    for key,label in [('decisiones','Decisiones'),('acuerdos','Acuerdos'),('pendientes','Pendientes'),('preguntas','Preguntas abiertas')]:
        items=summary.get(key) or []
        if items:
            d.add_heading(label,level=1)
            for item in items:
                if isinstance(item,dict):
                    t=str(item.get('texto') or item.get('text') or item)
                    extra=' · '.join(x for x in [str(item.get('responsable') or '').strip(),str(item.get('fecha') or '').strip()] if x)
                    d.add_paragraph(f'{t}{(" · "+extra) if extra else ""}',style='List Bullet')
                else:d.add_paragraph(str(item),style='List Bullet')
    d.add_heading('Transcripción',level=1)
    for row in rows:d.add_paragraph(f"[{row['time']}] {row['speaker']}: {row['text']}")
    d.save(path)

def build_pdf(meta,rows,summary,path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name='SkillBlue',parent=styles['Heading1'],textColor=colors.HexColor('#00416B')))
    story=[Paragraph('SKILLED PROYECTOS INDUSTRIALES',styles['SkillBlue']),Paragraph(meta['title'],styles['Title']),Spacer(1,8),Paragraph(f"Fecha: {meta['date']}<br/>Duración: {meta['duration']}<br/>Voces diferenciadas: {meta['speakers']}",styles['BodyText']),Spacer(1,10),Paragraph('Resumen ejecutivo',styles['SkillBlue']),Paragraph(str(summary.get('resumen') or 'Sin resumen.'),styles['BodyText'])]
    for key,label in [('decisiones','Decisiones'),('acuerdos','Acuerdos'),('pendientes','Pendientes'),('preguntas','Preguntas abiertas')]:
        items=summary.get(key) or []
        if items:
            story+=[Spacer(1,8),Paragraph(label,styles['SkillBlue'])]
            for item in items:story.append(Paragraph('• '+str(item.get('texto') if isinstance(item,dict) else item),styles['BodyText']))
    story+=[Spacer(1,10),Paragraph('Transcripción',styles['SkillBlue'])]
    for row in rows:story.append(Paragraph(f"[{row['time']}] <b>{row['speaker']}:</b> {row['text']}",styles['BodyText']))
    SimpleDocTemplate(str(path),pagesize=letter,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story)

def finalize_session(session):
    stamp=time.strftime('%Y%m%d-%H%M%S')
    slug=re.sub(r'[^A-Za-z0-9_-]+','_',session['device'])[:40]
    audio=OUT/f'{stamp}_{slug}.wav'; write_wav(audio,session['meeting'])
    diar=diarize(audio)
    segs=transcribe_pcm(session['meeting'],WHISPER_FINAL_MODEL,True)
    rows=rows_from_segments(segs,diar)
    summary=ollama_summary(rows)
    voices=len(set(r['speaker'] for r in rows)) if rows else 0
    meta={'title':'Minuta de reunión','date':time.strftime('%d/%m/%Y %H:%M'),'duration':time.strftime('%H:%M:%S',time.gmtime(len(session['meeting'])/2/SAMPLE_RATE)),'speakers':voices,'device':session['device']}
    base=OUT/f'{stamp}_{slug}'
    docx=Path(str(base)+'.docx'); pdf=Path(str(base)+'.pdf'); js=Path(str(base)+'.json')
    build_docx(meta,rows,summary,docx); build_pdf(meta,rows,summary,pdf)
    js.write_text(json.dumps({'meta':meta,'summary':summary,'transcript':rows},ensure_ascii=False,indent=2),encoding='utf-8')
    if not KEEP_AUDIO:audio.unlink(missing_ok=True)
    return {'docx':docx.name,'pdf':pdf.name,'json':js.name,'meta':meta}

@app.get('/health')
def health():return {'ok':True,'version':'136','whisper':WHISPER_MODEL,'finalModel':WHISPER_FINAL_MODEL,'pyannoteConfigured':bool(HF_TOKEN),'pyannoteModel':PYANNOTE_MODEL,'ollama':USE_OLLAMA,'smtpConfigured':bool(os.getenv('SMTP_HOST','') and (os.getenv('SMTP_FROM','') or os.getenv('SMTP_USER','')))}

@app.get('/files/{name}')
def file(name:str):
    path=(OUT/name).resolve()
    if path.parent!=OUT or not path.exists():raise HTTPException(404,'Archivo no encontrado.')
    return FileResponse(str(path),filename=path.name)

def smtp_send(to,subject,body,attachments=None):
    host=os.getenv('SMTP_HOST','');user=os.getenv('SMTP_USER','');password=os.getenv('SMTP_PASSWORD','');sender=os.getenv('SMTP_FROM',user);port=int(os.getenv('SMTP_PORT','587'));use_tls=os.getenv('SMTP_STARTTLS','true').lower() in ('1','true','yes','si','sí');use_ssl=os.getenv('SMTP_SSL','false').lower() in ('1','true','yes','si','sí')
    if not host or not sender:raise HTTPException(503,'SMTP no configurado. Completa SMTP_HOST y SMTP_FROM/SMTP_USER en el .env local de SKILL.')
    recipients=[str(x).strip() for x in to if str(x).strip()]
    if not recipients:raise HTTPException(400,'No hay destinatarios válidos.')
    msg=EmailMessage();msg['From']=sender;msg['To']=', '.join(recipients);msg['Subject']=subject;msg.set_content(body or '')
    for item in attachments or []:
        filename=str(item.get('filename') or 'archivo.bin')
        content=item.get('content') or b''
        content_type=str(item.get('contentType') or 'application/octet-stream')
        if isinstance(content,str):
            try:content=base64.b64decode(content,validate=True)
            except Exception:raise HTTPException(400,f'Adjunto inválido: {filename}')
        maintype,_,subtype=content_type.partition('/')
        if not subtype:maintype,subtype='application','octet-stream'
        msg.add_attachment(content,maintype=maintype,subtype=subtype,filename=filename)
    smtp_class=smtplib.SMTP_SSL if use_ssl else smtplib.SMTP
    with smtp_class(host,port,timeout=30) as smtp:
        if use_tls and not use_ssl:smtp.starttls()
        if user:smtp.login(user,password)
        smtp.send_message(msg)
    return {'ok':True,'from':sender,'to':recipients}

@app.post('/email')
def email(req:EmailRequest):
    attachments=[]
    for name in [req.docx,req.pdf]:
        if not name:continue
        path=(OUT/name).resolve()
        if path.parent!=OUT or not path.exists():continue
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document' if path.suffix=='.docx' else 'application/pdf'
        attachments.append({'filename':path.name,'content':path.read_bytes(),'contentType':content_type})
    return smtp_send(req.to,req.subject,req.body,attachments)

@app.post('/email-browser')
def email_browser(req:BrowserEmailRequest,x_skill_token:str|None=Header(default=None)):
    if TOKEN and TOKEN!='CAMBIA_ESTE_TOKEN' and str(x_skill_token or '')!=TOKEN:raise HTTPException(403,'Token inválido.')
    attachments=[{'filename':item.filename,'content':item.content,'contentType':item.contentType} for item in req.attachments]
    return smtp_send(req.to,req.subject,req.body,attachments)

@app.post('/process-browser')
async def process_browser(audio:UploadFile=File(...),title:str=Form('Reunión de trabajo'),participants:str=Form(''),x_skill_token:str|None=Header(default=None)):
    if TOKEN and TOKEN!='CAMBIA_ESTE_TOKEN' and str(x_skill_token or '')!=TOKEN:
        raise HTTPException(403,'Token inválido.')
    suffix=Path(audio.filename or 'reunion.webm').suffix or '.webm'
    fd,tmp=tempfile.mkstemp(suffix=suffix)
    os.close(fd);path=Path(tmp)
    try:
        with path.open('wb') as out:
            while True:
                chunk=await audio.read(1024*1024)
                if not chunk:break
                out.write(chunk)
        if path.stat().st_size<1024:raise HTTPException(400,'Audio vacío.')
        result=await asyncio.to_thread(finalize_uploaded_audio,path,title,participants)
        return result
    finally:
        path.unlink(missing_ok=True)

@app.websocket('/ws/meeting')
async def meeting_ws(ws:WebSocket):
    await ws.accept(); session={'device':'unknown','ring':deque(maxlen=SAMPLE_RATE*2*6),'meeting':bytearray(),'active':False,'lastCheck':0.0,'lastText':''}
    try:
        first=await ws.receive_text(); hello=json.loads(first)
        if str(hello.get('token') or '')!=TOKEN:await ws.close(code=4403);return
        session['device']=str(hello.get('deviceId') or 'skill')[:80]; clients[session['device']]=ws
        while True:
            msg=await ws.receive()
            data=msg.get('bytes')
            if data:
                session['ring'].extend(data)
                if session['active']:session['meeting'].extend(data)
                now=time.time()
                if now-session['lastCheck']>=2.2:
                    session['lastCheck']=now
                    ring=bytes(session['ring'])
                    try:text=await asyncio.to_thread(command_text,ring)
                    except Exception as exc:print('whisper:',exc);text=''
                    if text and text!=session['lastText']:
                        session['lastText']=text
                        if not session['active'] and is_start(text):
                            session['active']=True;session['meeting']=bytearray();await ws.send_text(json.dumps({'command':'START'}))
                        elif session['active'] and is_end(text):
                            session['active']=False;await ws.send_text(json.dumps({'command':'PROCESSING'}))
                            result=await asyncio.to_thread(finalize_session,session)
                            await ws.send_text(json.dumps({'command':'DONE','result':result},ensure_ascii=False));session['meeting']=bytearray()
    except WebSocketDisconnect:pass
    except Exception as exc:
        try:await ws.send_text(json.dumps({'command':'ERROR','message':str(exc)}))
        except Exception:pass
    finally:
        clients.pop(session['device'],None)
